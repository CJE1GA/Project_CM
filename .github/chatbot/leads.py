"""Lead storage for the chatbot, with cloud backends and local Excel fallback."""

import json
import os
import threading
import urllib.error
import urllib.request
from datetime import datetime

import re

from config import LEADS_FILE, LEAD_SOURCE

COLUMNS = ["Year", "Month", "Day", "Hour", "Name", "Phone Number", "Mail", "Birth Date", "Program", "Branch", "Source"]
GOOGLE_COLUMNS = ["Year", "Month", "Day", "Hour", "Name", "Phone Number", "Mail", "Birth Date", "Program", "Branch", "Source", "Identifier"]
UNSPECIFIED = "Sin especificar"
INSTAGRAM_HANDLE_RE = re.compile(r"@[_a-zA-Z0-9.]+")

_lock = threading.Lock()


def _path() -> str:
    return os.path.join(os.path.dirname(os.path.abspath(__file__)), LEADS_FILE)


def _google_config() -> dict:
    return {
        "spreadsheet_id": os.environ.get("GOOGLE_SHEETS_SPREADSHEET_ID", "").strip(),
        "worksheet_name": os.environ.get("GOOGLE_SHEETS_WORKSHEET_NAME", "Leads").strip() or "Leads",
        "service_account_json": os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON", "").strip(),
        "service_account_file": os.environ.get("GOOGLE_SERVICE_ACCOUNT_FILE", "").strip(),
    }


def _apps_script_config() -> dict:
    return {
        "url": os.environ.get("GOOGLE_APPS_SCRIPT_URL", "").strip(),
        "token": os.environ.get("GOOGLE_APPS_SCRIPT_TOKEN", "").strip(),
    }


def _column_letter(index: int) -> str:
    letters = ""
    while index > 0:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _split_timestamp(value) -> tuple[str, str, str, str]:
    if isinstance(value, datetime):
        return (
            str(value.year),
            f"{value.month:02d}",
            f"{value.day:02d}",
            f"{value.hour:02d}:00",
        )

    text = str(value or "").strip()
    if not text:
        return "", "", "", ""

    patterns = [
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d",
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%d/%m/%Y",
    ]
    for pattern in patterns:
        try:
            parsed = datetime.strptime(text, pattern)
            return (
                str(parsed.year),
                f"{parsed.month:02d}",
                f"{parsed.day:02d}",
                f"{parsed.hour:02d}:00",
            )
        except ValueError:
            continue

    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        return (
            str(parsed.year),
            f"{parsed.month:02d}",
            f"{parsed.day:02d}",
            f"{parsed.hour:02d}:00",
        )
    except ValueError:
        return text[:4], "", "", ""


def _sync_lead_to_apps_script(
    name: str,
    phone: str,
    email: str,
    birth_date: str,
    program: str,
    branch: str,
    identifier: str,
) -> bool:
    config = _apps_script_config()
    if not config["url"]:
        return False

    payload = {
        "token": config["token"],
        "action": "upsert_lead",
        "lead": {
            "year": str(datetime.now().year),
            "month": f"{datetime.now().month:02d}",
            "day": f"{datetime.now().day:02d}",
            "hour": f"{datetime.now().hour:02d}:00",
            "name": name,
            "phone": phone or "",
            "mail": email or "",
            "birth_date": birth_date or "",
            "program": program or UNSPECIFIED,
            "branch": branch or UNSPECIFIED,
            "source": LEAD_SOURCE,
            "identifier": identifier,
        },
    }
    request_obj = urllib.request.Request(
        config["url"],
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )

    try:
        with urllib.request.urlopen(request_obj, timeout=15) as response:
            response_body = response.read().decode("utf-8").strip()
        if not response_body:
            return True
        response_json = json.loads(response_body)
        return bool(response_json.get("ok"))
    except (urllib.error.URLError, urllib.error.HTTPError, TimeoutError, json.JSONDecodeError) as exc:
        print(f"[leads] No se pudo guardar en Google Apps Script: {exc}")
        return False


def _google_client():
    config = _google_config()
    if not config["spreadsheet_id"]:
        return None

    try:
        import gspread
    except ImportError:
        print("[leads] gspread no instalado; se usará Excel local.")
        return None

    try:
        if config["service_account_json"]:
            credentials_info = json.loads(config["service_account_json"])
            return gspread.service_account_from_dict(credentials_info)
        if config["service_account_file"]:
            return gspread.service_account(filename=config["service_account_file"])
    except Exception as exc:
        print(f"[leads] No se pudieron cargar credenciales de Google Sheets: {exc}")
        return None

    print("[leads] Faltan credenciales de Google Sheets; se usará Excel local.")
    return None


def _google_worksheet():
    config = _google_config()
    client = _google_client()
    if not client:
        return None

    try:
        spreadsheet = client.open_by_key(config["spreadsheet_id"])
        try:
            worksheet = spreadsheet.worksheet(config["worksheet_name"])
        except Exception:
            worksheet = spreadsheet.add_worksheet(title=config["worksheet_name"], rows=1000, cols=len(GOOGLE_COLUMNS))
        return worksheet
    except Exception as exc:
        print(f"[leads] No se pudo abrir Google Sheets: {exc}")
        return None


def _ensure_google_headers(worksheet):
    headers = worksheet.row_values(1)
    all_rows = worksheet.get_all_values()
    if not headers:
        worksheet.append_row(GOOGLE_COLUMNS)
        return
    if _google_sheet_needs_migration(headers, all_rows):
        _migrate_google_sheet(worksheet, headers, all_rows)


def _google_row_uses_legacy_schema(headers: list[str]) -> bool:
    normalized = [str(header or "").strip() for header in headers]
    return normalized[: min(len(normalized), 7)] == ["Date", "Name", "Phone Number", "Program", "Branch", "Source", "Identifier"][: min(len(normalized), 7)]


def _google_sheet_needs_migration(headers, all_rows) -> bool:
    if headers != GOOGLE_COLUMNS:
        return True
    if len(all_rows) <= 1:
        return False

    first_data_row = all_rows[1]
    if not first_data_row:
        return False

    first_cell = str(first_data_row[0] or "").strip()
    if not first_cell:
        return False

    if first_cell.isdigit() and len(first_cell) == 4:
        return False

    if re.match(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}$", first_cell):
        return True

    return len(first_data_row) < len(GOOGLE_COLUMNS) or _google_row_uses_legacy_schema(headers)


def _migrate_google_sheet(worksheet, headers, all_rows) -> None:
    if len(all_rows) <= 1:
        worksheet.clear()
        worksheet.append_row(GOOGLE_COLUMNS)
        return

    data_rows = all_rows[1:]
    migrated_rows = []
    is_legacy_schema = _google_row_uses_legacy_schema(headers)

    for row in data_rows:
        if not any(str(cell or "").strip() for cell in row):
            continue

        if is_legacy_schema:
            date_value = row[0] if len(row) > 0 else ""
            year, month, day, hour = _split_timestamp(date_value)
            migrated_rows.append([
                year,
                month,
                day,
                hour,
                row[1] if len(row) > 1 else "",
                row[2] if len(row) > 2 else "",
                "",
                "",
                row[3] if len(row) > 3 else UNSPECIFIED,
                row[4] if len(row) > 4 else UNSPECIFIED,
                row[5] if len(row) > 5 else LEAD_SOURCE,
                row[6] if len(row) > 6 else "",
            ])
            continue

        year, month, day, hour = _split_timestamp((row[0] + " " + row[1] + " " + row[2] + " " + row[3]).strip())
        migrated_rows.append([
            row[0] if len(row) > 0 else year,
            row[1] if len(row) > 1 else month,
            row[2] if len(row) > 2 else day,
            row[3] if len(row) > 3 else hour,
            row[4] if len(row) > 4 else "",
            row[5] if len(row) > 5 else "",
            row[6] if len(row) > 6 else "",
            row[7] if len(row) > 7 else "",
            row[8] if len(row) > 8 else UNSPECIFIED,
            row[9] if len(row) > 9 else UNSPECIFIED,
            row[10] if len(row) > 10 else LEAD_SOURCE,
            row[11] if len(row) > 11 else "",
        ])

    worksheet.clear()
    worksheet.append_row(GOOGLE_COLUMNS)
    if migrated_rows:
        worksheet.append_rows(migrated_rows, value_input_option="USER_ENTERED")


def _sync_lead_to_google_sheet(
    name: str,
    phone: str,
    email: str,
    birth_date: str,
    program: str,
    branch: str,
    identifier: str,
) -> bool:
    worksheet = _google_worksheet()
    if worksheet is None:
        return False

    try:
        _ensure_google_headers(worksheet)
        all_rows = worksheet.get_all_values()
        existing_row = None
        for index, row in enumerate(all_rows[1:], start=2):
            row_identifier = row[11].strip() if len(row) > 11 else ""
            if row_identifier == identifier:
                existing_row = index
                break

        values = [
            str(datetime.now().year),
            f"{datetime.now().month:02d}",
            f"{datetime.now().day:02d}",
            f"{datetime.now().hour:02d}:00",
            name,
            phone or "",
            email or "",
            birth_date or "",
            program or UNSPECIFIED,
            branch or UNSPECIFIED,
            LEAD_SOURCE,
            identifier,
        ]

        if existing_row is None:
            worksheet.append_row(values)
        else:
            worksheet.update(f"A{existing_row}:{_column_letter(len(GOOGLE_COLUMNS))}{existing_row}", [values])
        return True
    except Exception as exc:
        print(f"[leads] No se pudo guardar en Google Sheets: {exc}")
        return False


def _meaningful(value: str | None) -> bool:
    return bool(value and value.strip() and value.strip() != UNSPECIFIED)


def _coalesce(current: str | None, incoming: str | None) -> str:
    if _meaningful(incoming):
        return incoming.strip()
    if current and str(current).strip():
        return str(current).strip()
    if incoming and incoming.strip():
        return incoming.strip()
    return ""


def _legacy_workbook_row(sheet, row: int) -> bool:
    return not sheet.cell(row=row, column=11).value and sheet.cell(row=row, column=8).value == LEAD_SOURCE


def _migrate_legacy_workbook(sheet) -> None:
    headers = [sheet.cell(row=1, column=index + 1).value for index in range(len(COLUMNS))]
    if headers == COLUMNS:
        return

    legacy_rows = []
    for row in range(2, sheet.max_row + 1):
        year, month, day, hour = _split_timestamp(sheet.cell(row=row, column=1).value)
        legacy_rows.append([
            year,
            month,
            day,
            hour,
            sheet.cell(row=row, column=2).value,
            sheet.cell(row=row, column=3).value,
            sheet.cell(row=row, column=4).value,
            sheet.cell(row=row, column=5).value,
            sheet.cell(row=row, column=6).value,
            sheet.cell(row=row, column=7).value,
            sheet.cell(row=row, column=8).value,
        ])

    sheet.delete_rows(1, sheet.max_row)
    sheet.append(COLUMNS)
    for row_values in legacy_rows:
        sheet.append(row_values)


def _ensure_workbook():
    from openpyxl import Workbook, load_workbook

    path = _path()
    if os.path.exists(path):
        workbook = load_workbook(path)
        sheet = workbook.active
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Leads"
        sheet.append(COLUMNS)

    headers = [sheet.cell(row=1, column=index + 1).value for index in range(len(COLUMNS))]
    if headers != COLUMNS:
        _migrate_legacy_workbook(sheet)

    return workbook, sheet, path


def _extract_identifier(name: str | None) -> str:
    if not name:
        return ""
    match = INSTAGRAM_HANDLE_RE.search(str(name))
    return match.group(0) if match else str(name).strip()


def normalize_identifier(value: str | None) -> str:
    return _extract_identifier(value)


def upsert_lead(
    name: str,
    phone: str = "",
    email: str = "",
    birth_date: str = "",
    program: str = UNSPECIFIED,
    branch: str = UNSPECIFIED,
    identifier: str | None = None,
) -> None:
    with _lock:
        lookup_value = normalize_identifier(identifier or name)
        if _sync_lead_to_apps_script(name, phone, email, birth_date, program, branch, lookup_value):
            return
        if _sync_lead_to_google_sheet(name, phone, email, birth_date, program, branch, lookup_value):
            return

        try:
            from openpyxl import Workbook, load_workbook
        except ImportError:
            print("[leads] openpyxl no instalado; el lead no se guardó en Excel.")
            return

        del Workbook, load_workbook

        workbook, sheet, path = _ensure_workbook()
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        matching_rows = []

        for row in range(sheet.max_row, 1, -1):
            row_name = sheet.cell(row=row, column=2).value
            row_source = sheet.cell(row=row, column=8).value or sheet.cell(row=row, column=6).value
            row_identifier = _extract_identifier(row_name)
            if row_source == LEAD_SOURCE and row_identifier == lookup_value:
                matching_rows.append(row)

        if not matching_rows:
            sheet.append([
                str(datetime.now().year),
                f"{datetime.now().month:02d}",
                f"{datetime.now().day:02d}",
                f"{datetime.now().hour:02d}:00",
                name,
                phone or "",
                email or "",
                birth_date or "",
                program or UNSPECIFIED,
                branch or UNSPECIFIED,
                LEAD_SOURCE,
            ])
        else:
            merged_name = name
            merged_phone = phone or ""
            merged_email = email or ""
            merged_birth_date = birth_date or ""
            merged_program = program or UNSPECIFIED
            merged_branch = branch or UNSPECIFIED
            merged_year, merged_month, merged_day, merged_hour = _split_timestamp(datetime.now())

            for row in reversed(matching_rows):
                current_name = str(sheet.cell(row=row, column=5).value or "").strip()
                if len(current_name) > len(merged_name):
                    merged_name = current_name
                merged_phone = _coalesce(merged_phone, sheet.cell(row=row, column=6).value)
                merged_email = _coalesce(merged_email, sheet.cell(row=row, column=7).value)
                merged_birth_date = _coalesce(merged_birth_date, sheet.cell(row=row, column=8).value)
                if _legacy_workbook_row(sheet, row):
                    merged_program = _coalesce(merged_program, sheet.cell(row=row, column=6).value) or UNSPECIFIED
                    merged_branch = _coalesce(merged_branch, sheet.cell(row=row, column=7).value) or UNSPECIFIED
                else:
                    merged_program = _coalesce(merged_program, sheet.cell(row=row, column=9).value) or UNSPECIFIED
                    merged_branch = _coalesce(merged_branch, sheet.cell(row=row, column=10).value) or UNSPECIFIED

            for row in matching_rows:
                sheet.delete_rows(row, 1)

            sheet.append([
                merged_year,
                merged_month,
                merged_day,
                merged_hour,
                merged_name,
                merged_phone,
                merged_email,
                merged_birth_date,
                merged_program,
                merged_branch,
                LEAD_SOURCE,
            ])

        workbook.save(path)


def deduplicate_leads() -> int:
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError:
        print("[leads] openpyxl no instalado; no se pudo consolidar el Excel.")
        return 0

    del Workbook, load_workbook

    with _lock:
        workbook, sheet, path = _ensure_workbook()
        grouped = {}
        removed = 0

        for row in range(2, sheet.max_row + 1):
            year_value = sheet.cell(row=row, column=1).value
            month_value = sheet.cell(row=row, column=2).value
            day_value = sheet.cell(row=row, column=3).value
            hour_value = sheet.cell(row=row, column=4).value
            name_value = sheet.cell(row=row, column=5).value
            phone_value = sheet.cell(row=row, column=6).value
            if _legacy_workbook_row(sheet, row):
                mail_value = ""
                birth_date_value = ""
                program_value = sheet.cell(row=row, column=9).value
                branch_value = sheet.cell(row=row, column=10).value
                source_value = sheet.cell(row=row, column=11).value
            else:
                mail_value = sheet.cell(row=row, column=7).value
                birth_date_value = sheet.cell(row=row, column=8).value
                program_value = sheet.cell(row=row, column=9).value
                branch_value = sheet.cell(row=row, column=10).value
                source_value = sheet.cell(row=row, column=11).value

            if not name_value:
                continue

            normalized_name = str(name_value).strip()
            key = (_extract_identifier(normalized_name), str(source_value or "").strip())
            item = grouped.setdefault(
                key,
                {
                    "year": str(year_value or ""),
                    "month": str(month_value or ""),
                    "day": str(day_value or ""),
                    "hour": str(hour_value or ""),
                    "name": normalized_name,
                    "phone": "",
                    "mail": "",
                    "birth_date": "",
                    "program": UNSPECIFIED,
                    "branch": UNSPECIFIED,
                    "source": source_value or LEAD_SOURCE,
                    "count": 0,
                },
            )

            item["count"] += 1
            if len(normalized_name) > len(item["name"]):
                item["name"] = normalized_name
            item["phone"] = _coalesce(item["phone"], phone_value)
            item["mail"] = _coalesce(item["mail"], mail_value)
            item["birth_date"] = _coalesce(item["birth_date"], birth_date_value)
            item["program"] = _coalesce(item["program"], program_value) or UNSPECIFIED
            item["branch"] = _coalesce(item["branch"], branch_value) or UNSPECIFIED
            item["year"] = str(year_value or item["year"])
            item["month"] = str(month_value or item["month"])
            item["day"] = str(day_value or item["day"])
            item["hour"] = str(hour_value or item["hour"])

        rows = [COLUMNS]
        for item in grouped.values():
            removed += max(item["count"] - 1, 0)
            rows.append([
                item["year"],
                item["month"],
                item["day"],
                item["hour"],
                item["name"],
                item["phone"],
                item["mail"],
                item["birth_date"],
                item["program"],
                item["branch"],
                item["source"],
            ])

        sheet.delete_rows(1, sheet.max_row)
        for row in rows:
            sheet.append(row)

        workbook.save(path)
        return removed
